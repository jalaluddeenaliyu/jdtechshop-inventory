"""The Excel template people fill in to import products (see services/importer.py)."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from app.services.importer import COLUMNS

GREEN = "0E3B2E"
TEXT_COLUMNS = {"sku", "barcode", "imei1", "imei2", "serial", "model"}
EXAMPLES = [
    {"name": "Oraimo 20W charger", "tracking": "Quantity", "selling_price": 8500, "quantity": 30,
     "cost": 5200, "category": "Chargers", "brand": "Oraimo", "condition": "New"},
    {"name": "Tempered glass iPhone 15", "tracking": "Quantity", "selling_price": 2500,
     "quantity": 50, "cost": 900, "category": "Screen protectors"},
    {"name": "iPhone 13 128GB", "tracking": "Serial/IMEI", "selling_price": 650000,
     "cost": 560000, "category": "Phones", "brand": "Apple", "condition": "Used",
     "imei1": "356938035643809"},
    {"name": "iPhone 13 128GB", "tracking": "Serial/IMEI", "selling_price": 650000,
     "cost": 575000, "category": "Phones", "brand": "Apple", "condition": "Used",
     "imei1": "490154203237518"},
    {"name": "HP EliteBook 840 G5", "tracking": "Serial/IMEI", "selling_price": 420000,
     "cost": 350000, "category": "Laptops", "brand": "HP", "condition": "Refurbished",
     "serial": "5CG8251XYZ"},
]


def _sheet(wb: Workbook, title: str, rows: list[dict] | None, first: bool = False):
    ws = wb.active if first else wb.create_sheet(title)
    ws.title = title
    for c, (key, header, _, help_text) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=c, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=GREEN)
        cell.alignment = Alignment(vertical="center")
        width = max(14, len(header) + 4)
        if key == "name":
            width = 32
        elif key in ("imei1", "imei2", "specifications"):
            width = 20
        ws.column_dimensions[cell.column_letter].width = width
        for r in range(2, 1002 if first else 2 + len(rows or [])):
            if key in TEXT_COLUMNS:
                ws.cell(row=r, column=c).number_format = "@"  # keep IMEIs and codes exact
    ws.freeze_panes = "B2"
    for r, example in enumerate(rows or [], start=2):
        for c, (key, *_rest) in enumerate(COLUMNS, start=1):
            if key in example:
                ws.cell(row=r, column=c, value=example[key])
    return ws


def build_template() -> bytes:
    wb = Workbook()
    products = _sheet(wb, "Products", None, first=True)
    keys = [k for k, *_ in COLUMNS]

    def col(key: str) -> str:
        return products.cell(row=1, column=keys.index(key) + 1).column_letter

    for key, choices in (("tracking", '"Quantity,Serial/IMEI"'),
                         ("condition", '"New,Used,Refurbished"')):
        dv = DataValidation(type="list", formula1=choices, allow_blank=True)
        dv.add(f"{col(key)}2:{col(key)}1001")
        products.add_data_validation(dv)

    _sheet(wb, "Example", EXAMPLES)

    help_ws = wb.create_sheet("How to fill")
    help_ws.column_dimensions["A"].width = 22
    help_ws.column_dimensions["B"].width = 100
    lines = [
        ("JD TechShop — product import", ""),
        ("", ""),
        ("Where to type", "Fill in the “Products” sheet. Look at the “Example” sheet to see how."),
        ("Accessories", "One row per product, with Quantity (how many you have now) and "
                        "Cost price each."),
        ("Phones & laptops", "One row PER PHONE: the same product name on each row, with that "
                             "phone's IMEI (or serial number for laptops) and what it cost. "
                             "Leave Quantity blank."),
        ("Nothing in stock?", "Leave Quantity and Cost blank. The product is still created."),
        ("Money", "Plain numbers are best (8500). ₦8,500 also works."),
        ("Checking", "When you upload, the app shows a preview and lists any problems by row "
                     "number. Nothing is saved until every row is right."),
        ("", ""),
        ("Column", "What to put"),
    ] + [(header, help_text) for _, header, _, help_text in COLUMNS]
    for r, (a, b) in enumerate(lines, start=1):
        help_ws.cell(row=r, column=1, value=a)
        help_ws.cell(row=r, column=2, value=b).alignment = Alignment(wrap_text=True)
    help_ws["A1"].font = Font(bold=True, size=14, color=GREEN)
    help_ws["A10"].font = help_ws["B10"].font = Font(bold=True)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
