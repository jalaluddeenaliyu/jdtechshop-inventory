"""Importing products and opening stock from Excel / CSV."""
import io
from decimal import Decimal

from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from app.database import SessionLocal
from app.models import AuditLog, Product, ProductUnit, StockMovement

HEAD = ["Product name", "Counted by", "Selling price", "Quantity", "Cost price each", "Category",
        "Brand", "Condition", "IMEI 1", "IMEI 2", "Serial number", "SKU"]
GOOD_IMEIS = ["356938035643809", "490154203237518", "353918100000006"]


def xlsx(rows, head=HEAD) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    ws.append(head)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def upload(client, content: bytes, commit=False, name="stock.xlsx"):
    return client.post("/api/product-import", params={"commit": str(commit).lower()},
                       files={"file": (name, content, "application/octet-stream")})


def counts():
    with SessionLocal() as db:
        return (db.scalar(select(func.count()).select_from(Product)),
                db.scalar(select(func.count()).select_from(ProductUnit)),
                db.scalar(select(func.count()).select_from(StockMovement)))


MIXED = [
    ["Oraimo 20W charger", "Quantity", 8500, 30, 5200, "Chargers", "Oraimo", "New", None, None, None, None],
    ["Tempered glass iPhone 15", None, "₦2,500", "50", "900", "Screen protectors", None, None, None, None, None, None],
    ["iPhone 13 128GB", "Serial/IMEI", 650000, None, 560000, "Phones", "Apple", "Used", 356938035643809.0, None, None, None],
    ["iphone 13 128gb", "Serial/IMEI", 650000, None, 575000, "Phones", "Apple", "Used", GOOD_IMEIS[1], None, None, None],
    ["HP EliteBook 840 G5", None, 420000, None, 350000, "Laptops", "HP", "Refurbished", None, None, "5cg8251xyz", None],
    ["Lightning cable", None, 3000, None, None, None, None, None, None, None, None, None],
]


def test_preview_checks_everything_but_saves_nothing(owner_client):
    r = upload(owner_client, xlsx(MIXED))
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["saved"] is False and body["problems"] == []
    s = body["summary"]
    assert s["products"] == 5 and s["phones_and_laptops"] == 3 and s["items_in_stock"] == 83
    # 30×5200 + 50×900 + 560000 + 575000 + 350000
    assert Decimal(s["stock_value"]) == Decimal("1686000.00")
    assert sorted(s["new_categories"]) == ["Chargers", "Laptops", "Phones", "Screen protectors"]
    phone = next(p for p in body["products"] if p["name"] == "iPhone 13 128GB")
    assert phone["rows"] == [4, 5] and phone["quantity"] == 2
    assert Decimal(phone["cost_each"]) == Decimal("567500.00")
    assert counts() == (0, 0, 0)  # preview only


def test_import_saves_products_units_ledger_and_audit(owner_client):
    r = upload(owner_client, xlsx(MIXED), commit=True)
    assert r.status_code == 200 and r.json()["saved"] is True, r.text
    assert counts() == (5, 3, 5)  # 2 quantity openings + 3 phone/laptop units
    items = {p["name"]: p for p in owner_client.get("/api/products").json()["items"]}
    charger = items["Oraimo 20W charger"]
    assert charger["qty_on_hand"] == 30 and Decimal(charger["cost_each"]) == Decimal("5200.00")
    assert charger["category_name"] == "Chargers" and charger["brand_name"] == "Oraimo"
    assert items["Lightning cable"]["qty_on_hand"] == 0
    laptop = items["HP EliteBook 840 G5"]
    assert laptop["tracking_type"] == "serialized" and laptop["qty_on_hand"] == 1
    units = owner_client.get("/api/units", params={"q": "5CG8251XYZ"}).json()["items"]
    assert units[0]["serial_number"] == "5CG8251XYZ" and units[0]["condition"] == "refurbished"
    with SessionLocal() as db:
        assert db.scalar(select(func.count()).select_from(AuditLog)
                         .where(AuditLog.action == "products.imported")) == 1


def test_any_problem_means_nothing_is_saved_and_every_problem_is_listed(owner_client):
    rows = MIXED + [
        ["Bad price", None, "abc", 1, 100, None, None, None, None, None, None, None],
        ["No cost", None, 500, 4, None, None, None, None, None, None, None, None],
        ["Phone without cost", "Serial/IMEI", 100000, None, None, None, None, None, GOOD_IMEIS[2], None, None, None],
        ["Phone bad IMEI", None, 100000, None, 90000, None, None, None, "12345", None, None, None],
        ["Oraimo 20W charger", None, 8500, 3, 5000, None, None, None, None, None, None, None],
        [None, None, 100, 1, 1, None, None, None, None, None, None, None],
    ]
    r = upload(owner_client, xlsx(rows), commit=True)
    body = r.json()
    assert body["saved"] is False
    by_row = {p["rows"][0]: p["message"] for p in body["problems"]}
    assert "not a number" in by_row[8]
    assert "Cost price each is empty" in by_row[9]
    assert "cost you" in by_row[10]
    assert "IMEI" in by_row[11]
    assert "2 rows" in by_row[2] and "12" in by_row[2]  # duplicate charger rows 2 and 12
    assert "Product name is empty" in by_row[13]
    assert counts() == (0, 0, 0)


def test_imei_already_in_the_shop_or_twice_in_file_is_refused(owner_client):
    pid = owner_client.post("/api/products", json={
        "name": "Galaxy A15", "tracking_type": "serialized", "selling_price": "150000"}).json()["id"]
    owner_client.post(f"/api/products/{pid}/units/opening",
                      json={"units": [{"imei1": GOOD_IMEIS[0], "unit_cost": "120000"}]})
    rows = [
        ["iPhone 12", None, 400000, None, 300000, None, None, None, GOOD_IMEIS[0], None, None, None],
        ["iPhone 11", None, 300000, None, 200000, None, None, None, GOOD_IMEIS[1], None, None, None],
        ["iPhone 11", None, 300000, None, 200000, None, None, None, GOOD_IMEIS[1], None, None, None],
    ]
    body = upload(owner_client, xlsx(rows)).json()
    msgs = " | ".join(p["message"] for p in body["problems"])
    assert "already registered" in msgs and "Galaxy A15" in msgs
    assert "entered twice" in msgs


def test_existing_product_name_is_refused(owner_client):
    owner_client.post("/api/products", json={"name": "Lightning cable", "tracking_type": "quantity",
                                             "selling_price": "3000"})
    body = upload(owner_client, xlsx([MIXED[-1]])).json()
    assert "already exists" in body["problems"][0]["message"]


def test_csv_with_other_column_names_works(owner_client):
    csv_text = ("Item;Price;Qty;Cost;Brand\n"
                "Type-C cable 2m;4000;12;2100;Oraimo\n"
                "Car charger;5500;;;\n")
    r = upload(owner_client, csv_text.encode("utf-8-sig"), commit=True, name="stock.csv")
    assert r.status_code == 200 and r.json()["saved"], r.text
    items = {p["name"]: p for p in owner_client.get("/api/products").json()["items"]}
    assert items["Type-C cable 2m"]["qty_on_hand"] == 12 and items["Car charger"]["qty_on_hand"] == 0


def test_template_downloads_and_round_trips(owner_client):
    r = owner_client.get("/api/product-import/template")
    assert r.status_code == 200 and r.headers["content-type"].startswith(
        "application/vnd.openxmlformats")
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Products", "Example", "How to fill"]
    # The Example sheet, uploaded as it is, is a valid import.
    example = wb["Example"]
    out = Workbook()
    ws = out.active
    for row in example.iter_rows(values_only=True):
        ws.append(list(row))
    buf = io.BytesIO()
    out.save(buf)
    body = upload(owner_client, buf.getvalue()).json()
    assert body["problems"] == [] and body["summary"]["products"] == 4


def test_file_level_problems(owner_client):
    r = upload(owner_client, b"hello,world\n1,2\n", name="x.csv")
    assert r.status_code == 422 and "Product name" in r.json()["detail"]
    r = upload(owner_client, xlsx([], head=["Product name", "Selling price"]))
    assert r.status_code == 422 and "no products" in r.json()["detail"]
    r = upload(owner_client, b"\xd0\xcf\x11\xe0old", name="old.xls")
    assert r.status_code == 422 and ".xlsx" in r.json()["detail"]


def test_staff_cannot_import_or_download_template(staff_client):
    assert upload(staff_client, xlsx(MIXED)).status_code == 403
    assert staff_client.get("/api/product-import/template").status_code == 403
