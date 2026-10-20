"""Import products (and the stock already on the shelf) from an Excel or CSV file.

How a file becomes products:
- One row per accessory-type product (counted by quantity), with optional Quantity + Cost.
- One row per PHONE/LAPTOP UNIT for serial/IMEI products: rows that share a product name
  become one product, and each row with an IMEI or serial number becomes one unit at its
  own cost. A single row without an IMEI just creates the product (units come later).
- Missing categories and brands are created.

Everything runs in one database transaction using the same rules as the screens
(product checks, opening stock, IMEI checks, ledger, audit). Each product is tried inside its
own savepoint so ALL problems can be listed at once. If any row has a problem — or this is
only a preview — the whole transaction is rolled back: an import is all or nothing.
"""
from __future__ import annotations

import csv
import io
import re
from collections import OrderedDict
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from typing import Any

from fastapi import HTTPException, Request, status
from pydantic import ValidationError
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core import audit
from app.models import Brand, Category, Product, User
from app.schemas.catalog import ProductCreate
from app.schemas.units import UnitIn
from app.services import catalog as catalog_service
from app.services import units as unit_service

MAX_BYTES = 3 * 1024 * 1024
MAX_ROWS = 3000

# (key, template header, other headers people commonly use, help text for the template)
COLUMNS: list[tuple[str, str, tuple[str, ...], str]] = [
    ("name", "Product name", ("name", "product", "item", "item name", "product name"),
     "Required. One row per accessory. For phones/laptops: one row PER PHONE, same name."),
    ("tracking", "Counted by", ("type", "tracking", "counted", "tracking type"),
     "Quantity (chargers, cases…) or Serial/IMEI (phones, laptops). Blank = Quantity, "
     "or Serial/IMEI if the row has an IMEI/serial."),
    ("selling_price", "Selling price", ("price", "selling price (₦)", "sale price", "retail price"),
     "Required. Naira, e.g. 8500 or ₦8,500."),
    ("quantity", "Quantity", ("qty", "quantity in stock", "stock", "in stock", "how many",
                              "qty in stock"),
     "Quantity items only: how many you have now. Leave blank for phones (one row each)."),
    ("cost", "Cost price each", ("cost", "cost price", "cost each", "unit cost",
                                 "cost price (₦)", "buying price", "purchase price"),
     "What ONE cost you. Needed when Quantity is filled, and for every phone row."),
    ("category", "Category", (), "Optional. Created if it doesn't exist yet."),
    ("brand", "Brand", (), "Optional. Created if it doesn't exist yet."),
    ("model", "Model", (), "Optional."),
    ("sku", "SKU", ("code", "product code", "item code"),
     "Optional. Blank = one is made for you (JDT-00001…)."),
    ("barcode", "Barcode", ("ean", "upc"), "Optional."),
    ("condition", "Condition", (), "New, Used or Refurbished. Blank = New."),
    ("reorder_level", "Low-stock level", ("reorder level", "minimum stock", "min stock",
                                          "low stock level", "minimum"),
     "Optional. Blank = the shop default from Settings."),
    ("imei1", "IMEI 1", ("imei", "imei1"), "Phones: the 15-digit IMEI (dial *#06#)."),
    ("imei2", "IMEI 2", ("imei2",), "Dual-SIM phones: the second IMEI."),
    ("serial", "Serial number", ("serial", "serial no", "serial no.", "s/n", "sn"),
     "Laptops and other devices without an IMEI."),
    ("specifications", "Specifications", ("specs", "spec"), "Optional, e.g. 8GB RAM, 256GB."),
]
HEADER = {k: h for k, h, _, _ in COLUMNS}
FIELD_LABEL = {**HEADER, "unit_cost": HEADER["cost"], "serial_number": HEADER["serial"],
               "opening_quantity": HEADER["quantity"], "opening_unit_cost": HEADER["cost"],
               "tracking_type": HEADER["tracking"]}


def _norm_header(text: str) -> str:
    return re.sub(r"[^a-z0-9₦]+", " ", text.lower()).strip()


_ALIASES: dict[str, str] = {}
for _key, _header, _others, _ in COLUMNS:
    for _h in (_header, *_others):
        _ALIASES[_norm_header(_h)] = _key


# ---------------------------------------------------------------- reading the file

@dataclass
class Row:
    number: int  # the row number the person sees in Excel
    values: dict[str, str]


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        # Excel stores IMEIs, barcodes and whole prices as numbers: 356938035643809.0 → "356938035643809"
        return str(int(value)) if value.is_integer() else repr(value)
    return str(value).strip()


def read_rows(content: bytes, filename: str) -> list[Row]:
    if len(content) > MAX_BYTES:
        raise HTTPException(status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                            "That file is too big (3 MB at most). Split it into smaller files.")
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")) or content[:2] == b"PK":
        table = _read_xlsx(content)
    elif name.endswith((".csv", ".txt")) or not name:
        table = _read_csv(content)
    elif name.endswith(".xls"):
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY,
                            "Old .xls files can't be read. In Excel use File → Save As → "
                            "Excel Workbook (.xlsx), then upload that.")
    else:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY,
                            "Upload an Excel file (.xlsx) or a CSV file.")

    # The first row that names a "Product name" column is the header row.
    header_index, mapping = None, {}
    for i, cells in enumerate(table[:20]):
        found = {}
        for col, cell in enumerate(cells):
            key = _ALIASES.get(_norm_header(_cell_text(cell)))
            if key and key not in found.values():
                found[col] = key
        if "name" in found.values():
            header_index, mapping = i, found
            break
    if header_index is None:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY,
                            "Couldn't find a “Product name” column. Use the template, or make "
                            "sure the first row has the column names.")
    if "selling_price" not in mapping.values():
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY,
                            "Couldn't find a “Selling price” column.")

    rows: list[Row] = []
    for i, cells in enumerate(table[header_index + 1:], start=header_index + 2):
        values = {key: _cell_text(cells[col]) if col < len(cells) else ""
                  for col, key in mapping.items()}
        if not any(values.values()):
            continue  # blank line
        rows.append(Row(i, values))
    if not rows:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY,
                            "The file has column names but no products under them.")
    if len(rows) > MAX_ROWS:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY,
                            f"The file has {len(rows)} rows; {MAX_ROWS} at most per import. "
                            "Split it into smaller files.")
    return rows


def _read_xlsx(content: bytes) -> list[list[Any]]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY,
                            "That Excel file couldn't be opened. Save it again as .xlsx and retry.")
    sheet = wb["Products"] if "Products" in wb.sheetnames else wb.worksheets[0]
    table = [list(r) for r in sheet.iter_rows(values_only=True)]
    wb.close()
    return table


def _read_csv(content: bytes) -> list[list[Any]]:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "Couldn't read that CSV file.")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [row for row in csv.reader(io.StringIO(text), dialect)]


# ---------------------------------------------------------------- turning rows into products

class RowProblem(Exception):
    pass


def _money(text: str, label: str) -> str | None:
    t = re.sub(r"[₦,\s]|^NGN|^N(?=\d)", "", text.strip(), flags=re.IGNORECASE)
    if not t:
        return None
    try:
        value = Decimal(t)
    except InvalidOperation:
        raise RowProblem(f"{label} “{text}” is not a number.")
    if value < 0:
        raise RowProblem(f"{label} can't be negative.")
    return str(value)


def _whole(text: str, label: str) -> int | None:
    t = text.replace(",", "").strip()
    if not t:
        return None
    try:
        value = Decimal(t)
    except InvalidOperation:
        raise RowProblem(f"{label} “{text}” is not a number.")
    if value != value.to_integral_value():
        raise RowProblem(f"{label} must be a whole number.")
    return int(value)


def _tracking(text: str) -> str | None:
    t = _norm_header(text)
    if not t:
        return None
    if t in ("quantity", "qty", "q", "count", "counted", "accessory", "accessories"):
        return "quantity"
    if any(w in t for w in ("serial", "imei", "phone", "laptop", "unit")):
        return "serialized"
    raise RowProblem(f"{HEADER['tracking']} “{text}” should be Quantity or Serial/IMEI.")


def _condition(text: str) -> str | None:
    t = text.strip().lower()
    if not t:
        return None
    if t in ("new", "used", "refurbished"):
        return t
    if t.startswith("refurb"):
        return "refurbished"
    if t in ("uk used", "fairly used", "second hand", "secondhand", "tokunbo"):
        return "used"
    raise RowProblem(f"{HEADER['condition']} “{text}” should be New, Used or Refurbished.")


def _explain(e: Exception) -> str:
    if isinstance(e, ValidationError):
        err = e.errors()[0]
        loc = [str(p) for p in err.get("loc", []) if not isinstance(p, int)]
        label = FIELD_LABEL.get(loc[-1], loc[-1].replace("_", " ").capitalize()) if loc else ""
        msg = str(err.get("msg", "Invalid value")).removeprefix("Value error, ")
        return f"{label}: {msg}" if label else msg
    if isinstance(e, HTTPException):
        return str(e.detail)
    if isinstance(e, IntegrityError):
        return "An IMEI, serial number, SKU or barcode clashes with another product or unit."
    return str(e)


@dataclass
class Group:
    key: str
    rows: list[Row] = field(default_factory=list)


@dataclass
class Lookup:
    """Categories and brands by lower-case name, including ones created during this import."""
    model: type
    cache: dict[str, Any] = field(default_factory=dict)
    created: list[str] = field(default_factory=list)


def _resolve(db: Session, lookup: Lookup, name: str, actor: User, request: Request | None,
             kind: str) -> int | None:
    name = " ".join(name.split())
    if not name:
        return None
    key = name.lower()
    if key not in lookup.cache:
        found = db.scalar(select(lookup.model).where(func.lower(lookup.model.name) == key))
        if found is None:
            found = lookup.model(name=name[:80], is_active=True)
            db.add(found)
            db.flush()
            audit.record(db, action=f"{kind}.created", user=actor, entity_type=kind,
                         entity_id=found.id, details={"name": found.name, "via": "import"},
                         request=request)
            lookup.created.append(found.name)
        lookup.cache[key] = found
    obj = lookup.cache[key]
    if not obj.is_active:
        raise RowProblem(f"The {kind} “{obj.name}” is turned off. Turn it on in "
                         "Products → Categories & brands, or use another one.")
    return obj.id


def run_import(db: Session, content: bytes, filename: str, actor: User,
               request: Request | None, *, commit: bool) -> dict:
    rows = read_rows(content, filename)

    groups: "OrderedDict[str, Group]" = OrderedDict()
    blank_name_rows: list[int] = []
    for row in rows:
        name = " ".join(row.values.get("name", "").split())
        if not name:
            blank_name_rows.append(row.number)
            continue
        groups.setdefault(name.lower(), Group(name.lower())).rows.append(row)

    categories, brands = Lookup(Category), Lookup(Brand)
    results: list[dict] = []
    problems: list[dict] = [{"rows": [n], "product": None,
                             "message": f"{HEADER['name']} is empty."} for n in blank_name_rows]

    for group in groups.values():
        first = group.rows[0]
        v = first.values
        numbers = [r.number for r in group.rows]
        result = {"rows": numbers, "name": " ".join(v["name"].split()), "tracking_type": None,
                  "quantity": 0, "cost_each": None, "stock_value": "0.00",
                  "selling_price": None, "category": v.get("category") or None,
                  "brand": v.get("brand") or None, "error": None}
        savepoint = db.begin_nested()
        try:
            has_ids = [r for r in group.rows
                       if r.values.get("imei1") or r.values.get("imei2") or r.values.get("serial")]
            tracking = _tracking(v.get("tracking", "")) or ("serialized" if has_ids else "quantity")
            result["tracking_type"] = tracking
            if tracking == "quantity":
                if len(group.rows) > 1:
                    raise RowProblem(f"This product appears on {len(group.rows)} rows "
                                     f"({', '.join(map(str, numbers))}). Use one row with the "
                                     "total quantity, or give each a different name.")
                if has_ids:
                    raise RowProblem("This row has an IMEI/serial number, so “Counted by” "
                                     "should be Serial/IMEI.")
            price = _money(v.get("selling_price", ""), HEADER["selling_price"])
            if price is None:
                raise RowProblem(f"{HEADER['selling_price']} is empty.")
            result["selling_price"] = price
            body: dict[str, Any] = {
                "name": result["name"], "tracking_type": tracking, "selling_price": price,
                "sku": v.get("sku") or None, "barcode": v.get("barcode") or None,
                "model": v.get("model") or None,
                "specifications": v.get("specifications") or None,
                "reorder_level": _whole(v.get("reorder_level", ""), HEADER["reorder_level"]),
                "category_id": _resolve(db, categories, v.get("category", ""), actor, request,
                                        "category"),
                "brand_id": _resolve(db, brands, v.get("brand", ""), actor, request, "brand"),
            }
            if tracking == "quantity":
                body["condition"] = _condition(v.get("condition", "")) or "new"
                qty = _whole(v.get("quantity", ""), HEADER["quantity"])
                cost = _money(v.get("cost", ""), HEADER["cost"])
                if qty == 0:
                    qty = None if cost is None else qty
                if qty is not None or cost is not None:
                    if qty is None or qty < 1:
                        raise RowProblem(f"{HEADER['quantity']} is empty or 0 but a cost was "
                                         "given. Enter how many you have, or clear the cost.")
                    if cost is None:
                        raise RowProblem(f"{HEADER['cost']} is empty. Enter what one cost you.")
                    body["opening_quantity"], body["opening_unit_cost"] = qty, cost
            data = ProductCreate(**body)
            product = catalog_service.create_product(db, data, actor, request, commit=False)

            if tracking == "serialized":
                if len(group.rows) > 1 and len(has_ids) != len(group.rows):
                    missing = [r.number for r in group.rows if r not in has_ids]
                    raise RowProblem(f"Row {', '.join(map(str, missing))}: each phone row needs "
                                     f"its {HEADER['imei1']} or {HEADER['serial']}.")
                units: list[UnitIn] = []
                # Name the row only when this product spans several rows.
                at = (lambda r: f"Row {r.number}: ") if len(group.rows) > 1 else (lambda r: "")
                for r in has_ids:
                    rv = r.values
                    q = _whole(rv.get("quantity", ""), HEADER["quantity"])
                    if q not in (None, 1):
                        raise RowProblem(f"{at(r)}for phones use one row per phone "
                                         f"(with its IMEI) and leave {HEADER['quantity']} blank.")
                    cost = _money(rv.get("cost", ""), HEADER["cost"])
                    if cost is None:
                        raise RowProblem(f"{at(r)}{HEADER['cost']} is empty. Enter "
                                         "what this phone cost you.")
                    try:
                        units.append(UnitIn(
                            serial_number=rv.get("serial") or None, imei1=rv.get("imei1") or None,
                            imei2=rv.get("imei2") or None, unit_cost=cost,
                            condition=_condition(rv.get("condition", "")) or "new"))
                    except ValidationError as e:
                        raise RowProblem(f"{at(r)}{_explain(e)}")
                if units:
                    numbered = [r.number for r in has_ids]
                    unit_service.register_opening_units(
                        db, product, units, note="Imported from a spreadsheet", actor=actor,
                        request=request,
                        label=(lambda i: f"Row {numbered[i - 1]}: ") if len(group.rows) > 1
                        else (lambda i: ""))
            db.flush()
            db.refresh(product)
            value, each = catalog_service._valuation(
                product, catalog_service.unit_stock_values(db, [product]))
            result.update(quantity=product.qty_on_hand, stock_value=str(value),
                          cost_each=str(each) if each is not None else None, sku=product.sku)
            savepoint.commit()
        except (RowProblem, ValidationError, HTTPException, IntegrityError, ValueError) as e:
            savepoint.rollback()
            result["error"] = _explain(e)
            problems.append({"rows": numbers, "product": result["name"], "message": result["error"]})
        results.append(result)

    ok = [r for r in results if not r["error"]]
    summary = {
        "products": len(ok),
        "phones_and_laptops": sum(r["quantity"] for r in ok if r["tracking_type"] == "serialized"),
        "items_in_stock": sum(r["quantity"] for r in ok),
        "stock_value": str(sum((Decimal(r["stock_value"]) for r in ok), Decimal("0.00"))),
        "new_categories": categories.created,
        "new_brands": brands.created,
        "rows_read": len(rows),
    }
    problems.sort(key=lambda p: p["rows"][0])
    saved = commit and not problems
    if saved:
        audit.record(db, action="products.imported", user=actor, entity_type="import",
                     entity_id=None,
                     details={"file": (filename or "")[:120], **{k: v for k, v in summary.items()
                                                                 if not isinstance(v, list)}},
                     request=request)
        db.commit()
    else:
        db.rollback()
    return {"saved": saved, "summary": summary, "problems": problems, "products": results}
